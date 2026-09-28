from __future__ import annotations

import json
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

PROJECT_ROOT = Path(__file__).resolve().parents[2]
DATASET_ROOT = PROJECT_ROOT / "data" / "raw" / "bci_subjects_ept_v1"
OUTPUT_PATH = PROJECT_ROOT / "data" / "reports" / "session_modality_completeness.xlsx"

MODALITIES = (
    ("视频", "video"),
    ("EEG", "eeg"),
    ("PPG", "paired_ppg"),
    ("DOCX", "annotated_transcript"),
)


def load_rows() -> list[dict[str, object]]:
    rows: list[dict[str, object]] = []
    for session_dir in sorted(DATASET_ROOT.glob("session_[0-9][0-9][0-9]")):
        metadata_path = session_dir / "session_metadata.json"
        metadata = json.loads(metadata_path.read_text(encoding="utf-8"))
        available = metadata.get("available_sources", {})
        rows.append(
            {
                "session": session_dir.name,
                **{key: bool(available.get(key, False)) for _, key in MODALITIES},
            }
        )
    return rows


def category(row: dict[str, object]) -> str:
    flags = [bool(row[key]) for _, key in MODALITIES]
    if all(flags):
        return "四类齐全"
    if row["video"] and not row["eeg"] and not row["paired_ppg"] and not row["annotated_transcript"]:
        return "仅视频（另有派生音频）"
    if row["annotated_transcript"] and not row["video"] and not row["eeg"] and not row["paired_ppg"]:
        return "仅DOCX（按本表四类）"
    return "部分模态"


def build_workbook(rows: list[dict[str, object]]) -> Workbook:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Session完整度"

    headers = ["Session", *[label for label, _ in MODALITIES], "本表分类"]
    sheet.append(headers)

    for row in rows:
        sheet.append(
            [
                row["session"],
                *["✓" if row[key] else "" for _, key in MODALITIES],
                category(row),
            ]
        )

    total_row = len(rows) + 2
    sheet.cell(total_row, 1, "合计")
    for col in range(2, 2 + len(MODALITIES)):
        letter = get_column_letter(col)
        sheet.cell(total_row, col, f'=COUNTIF({letter}2:{letter}{total_row - 1},"✓")')
    sheet.cell(total_row, 6, f"共 {len(rows)} 个 session")

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    full_fill = PatternFill("solid", fgColor="E2F0D9")
    video_only_fill = PatternFill("solid", fgColor="FCE4D6")
    docx_only_fill = PatternFill("solid", fgColor="FFF2CC")
    partial_fill = PatternFill("solid", fgColor="DDEBF7")
    total_fill = PatternFill("solid", fgColor="D9EAD3")
    thin = Side(style="thin", color="D9E1F2")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    fill_by_category = {
        "四类齐全": full_fill,
        "仅视频（另有派生音频）": video_only_fill,
        "仅DOCX（按本表四类）": docx_only_fill,
        "部分模态": partial_fill,
    }
    for row_index in range(2, total_row):
        row_category = sheet.cell(row_index, 6).value
        row_fill = fill_by_category[row_category]
        for cell in sheet[row_index]:
            cell.fill = row_fill
            cell.border = border
            cell.alignment = Alignment(horizontal="center", vertical="center")

    for cell in sheet[total_row]:
        cell.fill = total_fill
        cell.font = Font(bold=True)
        cell.border = border
        cell.alignment = Alignment(horizontal="center", vertical="center")

    widths = {"A": 18, "B": 11, "C": 11, "D": 11, "E": 11, "F": 28}
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    sheet.row_dimensions[1].height = 25
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:F{total_row - 1}"
    sheet.sheet_view.showGridLines = False

    notes = workbook.create_sheet("说明")
    notes.append(["项目", "说明"])
    notes.append(["统计范围", f"{len(rows)} 个 session（session_001 至 session_041）"])
    notes.append(["✓", "当前 session_metadata.json 标记该模态存在"])
    notes.append(["空白", "当前元数据未登记该模态"])
    notes.append(["视频", "原始 MP4；由视频抽取的 WAV 未单独列出"])
    notes.append(["PPG", "paired_ppg，即配对 PPG 信号"])
    notes.append(["DOCX", "annotated_transcript，即标注文本来源"])
    notes.append(["颜色", "绿色=四类齐全；橙色=仅视频；黄色=仅DOCX；蓝色=部分模态"])
    for cell in notes[1]:
        cell.fill = header_fill
        cell.font = header_font
    for row in notes.iter_rows():
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    notes.column_dimensions["A"].width = 18
    notes.column_dimensions["B"].width = 72
    notes.freeze_panes = "A2"
    notes.sheet_view.showGridLines = False

    return workbook


def main() -> None:
    rows = load_rows()
    if len(rows) != 41:
        raise RuntimeError(f"Expected 41 sessions, found {len(rows)}")

    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    workbook = build_workbook(rows)
    workbook.save(OUTPUT_PATH)

    check = load_workbook(OUTPUT_PATH, data_only=False, read_only=True)
    sheet = check["Session完整度"]
    if sheet.max_row != 43 or sheet.max_column != 6:
        raise RuntimeError(
            f"Unexpected workbook dimensions: {sheet.max_row} rows x {sheet.max_column} columns"
        )
    check.close()
    print(OUTPUT_PATH)


if __name__ == "__main__":
    main()
