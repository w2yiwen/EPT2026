"""Data entry point for extracting staged BCI DOCX text and core properties."""

from __future__ import annotations

import argparse
import json
from pathlib import Path

from docx import Document


def extract_docx(path: Path) -> dict[str, object]:
    try:
        document = Document(path)
    except Exception as error:  # Preserve malformed legacy files in the audit.
        return {
            "path": str(path),
            "name": path.name,
            "bytes": path.stat().st_size,
            "parse_error": f"{type(error).__name__}: {error}",
            "paragraphs": [],
            "tables": [],
            "text": "",
        }
    paragraphs = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
    tables = [
        [[cell.text.strip() for cell in row.cells] for row in table.rows]
        for table in document.tables
    ]
    core = document.core_properties
    table_text = [cell for table in tables for row in table for cell in row if cell]
    return {
        "path": str(path),
        "name": path.name,
        "bytes": path.stat().st_size,
        "created": core.created.isoformat() if core.created else None,
        "modified": core.modified.isoformat() if core.modified else None,
        "title": core.title,
        "subject": core.subject,
        "author": core.author,
        "paragraphs": paragraphs,
        "tables": tables,
        "text": "\n".join([*paragraphs, *table_text]),
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--dataset-root", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()

    dataset_root = args.dataset_root.resolve(strict=True)
    sessions: list[dict[str, object]] = []
    for session_root in sorted(dataset_root.glob("session_*")):
        if not session_root.is_dir():
            continue
        metadata_path = session_root / "session_metadata.json"
        metadata = json.loads(metadata_path.read_text(encoding="utf-8")) if metadata_path.is_file() else {}
        documents = [extract_docx(path) for path in sorted(session_root.glob("*.docx"))]
        sessions.append(
            {
                "session_id": metadata.get("session_id", session_root.name),
                "subject_id": metadata.get("subject_id"),
                "documents": documents,
            }
        )

    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(
        json.dumps({"dataset_root": str(dataset_root), "sessions": sessions}, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )


if __name__ == "__main__":
    main()
